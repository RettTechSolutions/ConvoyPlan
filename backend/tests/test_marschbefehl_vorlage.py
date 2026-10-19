"""Marschbefehl aus der Word-Vorlage der Organisation.

Die Zusagen:

- Die Beispielvorlage kennt jedes Feld, das der Export füllt, und keins mehr;
  sie besteht die eigene Prüfung, und ihre Feldübersicht und Steuerzeilen
  landen nicht im fertigen Dokument.
- Was in einer Vorlage nicht stimmt — Tippfehler im Feld, kaputter
  Platzhalter, Makros, keine .docx —, fällt beim Hochladen auf, mit einer
  Meldung für den Org-Admin, und wird nicht gespeichert.
- Eine Vorlage ist fremder Code: Ausdrücke, die Speicher oder Rechenzeit
  sprengen oder ins Innere von Python greifen, werden abgewiesen — und was
  die Sandbox nicht erkennt, bricht der Kindprozess nach Frist oder am
  Speicherdeckel ab, statt den Server zu blockieren.
- Freitexte gehen unverändert hinein: Sonderzeichen machen das Dokument nicht
  kaputt, Zeilenumbrüche bleiben Umbrüche.
- Hochladen und Entfernen darf nur der Org-Admin; jedes Mitglied bekommt die
  Beispielvorlage. Der Export nimmt die eigene Vorlage, sonst das Beispiel.

Zuerst die Regeln ohne Datenbank, danach der Weg durch die App.
"""
import io
import zipfile
from types import SimpleNamespace

import pytest
from docx import Document
from sqlalchemy import delete

from app.database import AsyncSessionLocal
from app.models.marschbefehl_vorlage import MarschbefehlVorlage
from app.services import marschbefehl_vorlage as mv
from tests.aktionsseite_fixtures import (  # noqa: F401 — Fixtures, per Import aktiviert
    GEHEIMER_KONVOINAME,
    RUFNAME,
    aktion,
    client,
    h,
    reset_db_engine,
)


def _vorlage(*absaetze: str) -> bytes:
    doc = Document()
    for a in absaetze:
        doc.add_paragraph(a)
    puffer = io.BytesIO()
    doc.save(puffer)
    return puffer.getvalue()


def _text(docx: bytes) -> str:
    doc = Document(io.BytesIO(docx))
    teile = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            teile.extend(c.text for c in r.cells)
    return "\n".join(teile)


def _leerer_konvoi(**felder):
    werte = dict(
        name="Leer", organization=None, start_time=None, lage=None, auftrag=None, marschform=None,
        ablaufpunkt=None, ablaufzeit=None, ablaufführer=None, speed_urban_kmh=40, speed_rural_kmh=65,
        spacing_urban_m=15, spacing_rural_m=50, spacing_motorway_m=100, versorgung=None,
        funkgruppe=None, anlagen=None,
    )
    werte.update(felder)
    return SimpleNamespace(**werte)


# ── Felder ───────────────────────────────────────────────────────────────


def test_feldkatalog_deckt_den_kontext():
    """Jedes Feld, das der Export füllt, steht in der Übersicht — und keins mehr.
    Sonst fände der Org-Admin ein Feld nicht oder bekäme beim Hochladen
    „unbekanntes Feld“ für etwas, das die Übersicht anbietet."""
    k = mv.beispiel_kontext()
    gruppen = {name: {f.name for f in felder} for name, _, felder in mv.GRUPPEN}
    listen = {li.name: {f.name for f in li.felder} for li in mv.LISTEN}
    assert set(k) == set(gruppen) | set(listen)
    for name, felder in gruppen.items():
        assert set(k[name]) == felder, name
    for name, felder in listen.items():
        assert k[name], f"Beispiel ohne Eintrag in {name} — der Probelauf prüfte die Schleife nicht"
        for eintrag in k[name]:
            assert set(eintrag) == felder, name


def test_werte_sind_texte():
    """Fertig formatiert, damit in Word niemand Zahlen formatieren muss."""
    k = mv.beispiel_kontext()
    for name, _, _ in mv.GRUPPEN:
        assert all(isinstance(v, str) for v in k[name].values()), name
    for li in mv.LISTEN:
        assert all(isinstance(v, str) for e in k[li.name] for v in e.values()), li.name


def test_deutsche_formate():
    k = mv.beispiel_kontext()
    assert k["konvoi"]["marschstrecke"] == "412,3 km"
    assert k["konvoi"]["datum"] == "07.11.2026"
    assert k["fahrzeuge"][1]["gewicht"] == "18.000 kg"
    assert k["fahrzeuge"][1]["hoehe"] == "3,40 m"
    assert k["kanalwechsel"][0]["km"] == "87,4 km"


def test_leere_felder_sind_leer_nicht_strich():
    """``""`` statt „-“, damit ``{% if konvoi.lage %}`` greift."""
    k = mv.kontext(_leerer_konvoi(), [], [], None)
    assert k["konvoi"]["lage"] == ""
    assert k["konvoi"]["marschziel"] == ""
    assert k["konvoi"]["marschstrecke"] == ""
    assert k["fahrzeuge"] == [] and k["kanalwechsel"] == []


# ── Beispielvorlage ──────────────────────────────────────────────────────


def test_beispielvorlage_besteht_die_eigene_pruefung():
    mv.pruefe(mv.beispielvorlage())


def test_beispielvorlage_fuellt_und_verschluckt_steuerzeilen():
    text = _text(mv.fuellen(mv.beispielvorlage(), mv.beispiel_kontext()))
    assert "Hilfskonvoi Ahrtal" in text
    assert "Florian Musterstadt 66/1" in text  # Zeile je Fahrzeug
    assert "ILS Bayerischer Untermain" in text
    assert "{{" not in text and "{%" not in text
    assert "Feldübersicht" not in text


def test_beispielvorlage_nennt_jedes_feld():
    vorlage = _text(mv.beispielvorlage())
    for name, _, felder in mv.GRUPPEN:
        for f in felder:
            assert f"{{{{ {name}.{f.name} }}}}" in vorlage
    for li in mv.LISTEN:
        for f in li.felder:
            assert f"{{{{ {li.name[0]}.{f.name} }}}}" in vorlage


def test_leere_abschnitte_fallen_weg():
    text = _text(mv.fuellen(mv.beispielvorlage(), mv.kontext(_leerer_konvoi(), [], [], None)))
    assert "1. Lage" not in text
    assert "Kanalwechsel" not in text
    assert "Durchfahrtshöhen" not in text
    assert "3. Durchführung" in text


# ── Werte im Dokument ────────────────────────────────────────────────────


def test_sonderzeichen_und_zeilenumbrueche():
    lage = 'Brücke A & B gesperrt <Umleitung> "U 12"\nzweite Zeile\x07\x0c'
    k = mv.kontext(_leerer_konvoi(lage=lage), [], [], None)
    aus = mv.fuellen(_vorlage("{{ konvoi.lage }}"), k)
    text = _text(aus)
    assert 'Brücke A & B gesperrt <Umleitung> "U 12"\nzweite Zeile' in text
    # Steuerzeichen hätten docxtpl einen Absatz- bzw. Seitenumbruch machen lassen.
    assert "\x07" not in text and "\x0c" not in text


def test_platzhalter_in_einem_wert_wird_nicht_ausgewertet():
    k = mv.kontext(_leerer_konvoi(name="{{ 7 * 7 }}"), [], [], None)
    assert "{{ 7 * 7 }}" in _text(mv.fuellen(_vorlage("{{ konvoi.name }}"), k))


# ── Prüfen beim Hochladen ────────────────────────────────────────────────


def test_tippfehler_im_feld_wird_gemeldet():
    with pytest.raises(mv.VorlageFehler, match="nmae"):
        mv.pruefe(_vorlage("{{ konvoi.nmae }}"))


def test_unbekannte_liste_wird_gemeldet():
    with pytest.raises(mv.VorlageFehler, match="fahrzuege"):
        mv.pruefe(_vorlage("{% for f in fahrzuege %}{{ f.name }}{% endfor %}"))


def test_kaputter_platzhalter_wird_gemeldet():
    with pytest.raises(mv.VorlageFehler, match="Platzhalter"):
        mv.pruefe(_vorlage("{{ konvoi.name }"))


@pytest.mark.parametrize("ausdruck", [
    '{{ "x" * 100000000 }}',
    "{{ 999 ** 999 }}",
    "{% for i in range(1000) %}x{% endfor %}",
    '{{ "{:>999999999}".format(1) }}',
    '{{ "{:>999999999}"["format"](1) }}',
    '{{ "%0999999999d" % 1 }}',
    '{{ "x"|center(999999999) }}',
    "{{ konvoi.__class__.__mro__ }}",
    "{{ konvoi.name.__class__ }}",
    "{{ lipsum.__globals__ }}",
])
def test_sandbox_weist_ab(ausdruck):
    with pytest.raises(mv.VorlageFehler):
        mv.pruefe(_vorlage(ausdruck))


def test_rechenzeit_bricht_die_frist_ab(monkeypatch):
    """Zwei Selbstaufrufe je Ebene, 2**40 Aufrufe — keine Sandbox sieht das."""
    monkeypatch.setattr(mv, "FRIST_S", 3.0)
    bombe = "{% macro m(n) %}{% if n %}{{ m(n-1) }}{{ m(n-1) }}{% endif %}{% endmacro %}{{ m(40) }}"
    with pytest.raises(mv.VorlageFehler, match="zu lange"):
        mv.pruefe(_vorlage(bombe))


def test_speicher_bricht_am_deckel_ab():
    bombe = (
        "{% set ns = namespace(s='x' * 9000) %}"
        "{% for a in konvoi.lage %}{% for b in konvoi.lage %}{% set ns.s = ns.s ~ ns.s %}"
        "{% endfor %}{% endfor %}{{ ns.s|length }}"
    )
    with pytest.raises(mv.VorlageFehler, match="Speicher"):
        mv.pruefe(_vorlage(bombe))


def test_erlaubte_ausdruecke():
    mv.pruefe(_vorlage(
        "{{ konvoi.funkgruppe or '-' }}",
        "{{ konvoi.name|upper }}",
        "{% if fahrzeuge|length > 1 %}{{ fahrzeuge|map(attribute='name')|join(', ') }}{% endif %}",
        "{{ '=' * 20 }}",
    ))


def test_keine_docx():
    with pytest.raises(mv.VorlageFehler, match="keine .docx"):
        mv.pruefe(b"%PDF-1.7 nicht wirklich")


def test_zu_gross():
    with pytest.raises(mv.VorlageFehler, match="zu groß"):
        mv.pruefe(b"0" * (mv.MAX_BYTES + 1))


def _umschreiben(docx: bytes, aendern) -> bytes:
    rein = zipfile.ZipFile(io.BytesIO(docx))
    puffer = io.BytesIO()
    with zipfile.ZipFile(puffer, "w") as raus:
        for e in rein.infolist():
            raus.writestr(e, aendern(e.filename, rein.read(e)))
        if aendern("word/vbaProject.bin", None) is not None:
            raus.writestr("word/vbaProject.bin", b"\0")
    return puffer.getvalue()


def test_makros_werden_abgewiesen():
    mit_makro = _umschreiben(_vorlage("x"), lambda n, d: b"" if n == "word/vbaProject.bin" else d)
    with pytest.raises(mv.VorlageFehler, match="Makros"):
        mv.pruefe(mit_makro)


def test_dokumentvorlage_dotx_wird_abgewiesen():
    def zu_dotx(name, daten):
        if name == "[Content_Types].xml":
            return daten.replace(b"document.main+xml", b"template.main+xml")
        return daten
    with pytest.raises(mv.VorlageFehler, match=".dotx"):
        mv.pruefe(_umschreiben(_vorlage("x"), zu_dotx))


def test_zip_bombe_wird_abgewiesen():
    def aufblasen(name, daten):
        if name == "word/document.xml":
            return daten + b" " * (mv.MAX_ENTPACKT + 1)
        return daten
    bombe = _umschreiben(_vorlage("x"), aufblasen)
    assert len(bombe) < mv.MAX_BYTES
    with pytest.raises(mv.VorlageFehler, match="entpackt"):
        mv.pruefe(bombe)


# ── Durch die App ────────────────────────────────────────────────────────


@pytest.fixture
async def ohne_vorlage(aktion):
    yield aktion
    async with AsyncSessionLocal() as db:
        await db.execute(delete(MarschbefehlVorlage).where(MarschbefehlVorlage.organization_id == aktion.org_id))
        await db.commit()


def _datei(inhalt: bytes, name: str = "Unser Marschbefehl.docx") -> dict:
    return {"file": (name, inhalt, mv.DOCX_MIME)}


async def test_mitglied_bekommt_beispiel_aber_laedt_nicht_hoch(client, ohne_vorlage):
    a = ohne_vorlage
    r = await client.get("/api/org/marschbefehl-vorlage/beispiel", headers=h(a.fahrer))
    assert r.status_code == 200
    assert r.headers["content-type"] == mv.DOCX_MIME
    mv.pruefe(r.content)

    r = await client.get("/api/org/marschbefehl-vorlage", headers=h(a.fahrer))
    assert r.json() == {"eigene": False, "dateiname": None, "groesse": None, "hochgeladen_at": None}

    r = await client.get("/api/org/marschbefehl-vorlage/felder", headers=h(a.fahrer))
    assert {g["name"] for g in r.json()["gruppen"]} == {"konvoi", "organisation", "erstellt"}

    r = await client.post("/api/org/marschbefehl-vorlage", files=_datei(_vorlage("x")), headers=h(a.fahrer))
    assert r.status_code == 403
    r = await client.delete("/api/org/marschbefehl-vorlage", headers=h(a.fahrer))
    assert r.status_code == 403


async def test_eigene_vorlage_hochladen_exportieren_entfernen(client, ohne_vorlage):
    a = ohne_vorlage
    # Ohne eigene Vorlage: Export aus dem Beispiel.
    r = await client.get(f"/api/convoys/{a.convoy_id}/export/docx", headers=h(a.admin))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == mv.DOCX_MIME
    text = _text(r.content)
    assert "MARSCHBEFEHL" in text and GEHEIMER_KONVOINAME in text and RUFNAME in text

    eigene = _vorlage(
        "Einsatzbefehl der Wehr für {{ konvoi.name }}",
        "{% for f in fahrzeuge %}{{ f.pos }}: {{ f.funkrufname }}; {% endfor %}",
    )
    r = await client.post("/api/org/marschbefehl-vorlage", files=_datei(eigene), headers=h(a.admin))
    assert r.status_code == 200, r.text
    assert r.json()["eigene"] is True
    assert r.json()["dateiname"] == "Unser Marschbefehl.docx"

    r = await client.get(f"/api/convoys/{a.convoy_id}/export/docx", headers=h(a.admin))
    text = _text(r.content)
    assert f"Einsatzbefehl der Wehr für {GEHEIMER_KONVOINAME}" in text
    assert f"1: {RUFNAME}; 2: Florian Landshut 11/12;" in text

    r = await client.get("/api/org/marschbefehl-vorlage/datei", headers=h(a.admin))
    assert r.content == eigene

    r = await client.delete("/api/org/marschbefehl-vorlage", headers=h(a.admin))
    assert r.json()["eigene"] is False
    r = await client.get(f"/api/convoys/{a.convoy_id}/export/docx", headers=h(a.admin))
    assert "Einsatzbefehl" not in _text(r.content)
    r = await client.get("/api/org/marschbefehl-vorlage/datei", headers=h(a.admin))
    assert r.status_code == 404


async def test_fehlerhafte_vorlage_wird_nicht_gespeichert(client, ohne_vorlage):
    a = ohne_vorlage
    r = await client.post(
        "/api/org/marschbefehl-vorlage", files=_datei(_vorlage("{{ konvoi.nmae }}")), headers=h(a.admin)
    )
    assert r.status_code == 400
    assert "nmae" in r.json()["detail"]
    r = await client.get("/api/org/marschbefehl-vorlage", headers=h(a.admin))
    assert r.json()["eigene"] is False


async def test_vorlage_geht_als_datei_in_den_organisations_export(client, ohne_vorlage):
    from app.models.organization import Organization
    from app.services import org_export

    a = ohne_vorlage
    eigene = _vorlage("{{ konvoi.name }}")
    r = await client.post("/api/org/marschbefehl-vorlage", files=_datei(eigene), headers=h(a.admin))
    assert r.status_code == 200, r.text
    async with AsyncSessionLocal() as db:
        org = await db.get(Organization, a.org_id)
        bundle = await org_export.build_org_export(db, org)
        vorlage = await org_export.marschbefehl_vorlage(db, org)
    assert bundle["marschbefehl_vorlage"]["dateiname"] == "Unser Marschbefehl.docx"
    assert "daten" not in bundle["marschbefehl_vorlage"]
    zf = zipfile.ZipFile(io.BytesIO(org_export.export_zip(bundle, vorlage=vorlage)))
    assert zf.read("vorlagen/marschbefehl.docx") == eigene
